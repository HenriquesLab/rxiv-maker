"""Equations export as editable Office Math, which journals ask for."""

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from rxiv_maker.exporters.docx_writer import DocxWriter

SCHRODINGER = r"i\hbar\frac{\partial}{\partial t}\Psi(\mathbf{r},t) = \hat{H}\Psi(\mathbf{r},t)"


def _writer():
    return DocxWriter.__new__(DocxWriter)


def test_display_equation_becomes_centred_office_math():
    paragraph = Document().add_paragraph()

    assert _writer()._add_display_equation_omml(paragraph, SCHRODINGER) is True

    wrapper = paragraph._element.find(qn("m:oMathPara"))
    assert wrapper is not None
    assert wrapper.find(qn("m:oMathParaPr")).find(qn("m:jc")).get(qn("m:val")) == "center"

    math = etree.tostring(wrapper, encoding="unicode")
    assert "<m:f>" in math, "fraction lost"
    assert "<m:acc>" in math, "hat accent lost"
    assert "\u2202" in math and "\u03a8" in math


def test_inline_equation_keeps_its_superscript():
    paragraph = Document().add_paragraph()

    _writer()._add_inline_equation(paragraph, "E = mc^2")

    math = etree.tostring(paragraph._element, encoding="unicode")
    assert "<m:sSup>" in math, "superscript flattened to text"


def test_flat_conversion_defers_to_the_image_path():
    paragraph = Document().add_paragraph()

    assert _writer()._add_display_equation_omml(paragraph, "x") is False
    assert paragraph._element.find(qn("m:oMathPara")) is None


def test_unconvertible_latex_defers_to_the_image_path():
    paragraph = Document().add_paragraph()

    assert _writer()._add_display_equation_omml(paragraph, "") is False
