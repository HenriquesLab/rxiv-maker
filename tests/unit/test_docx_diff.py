"""Marked revisions keep the formatting of the runs they mark."""

from docx import Document

from rxiv_maker.exporters.docx_diff import (
    DELETED_COLOR,
    INSERTED_HIGHLIGHT,
    _mark_paragraph,
    _mark_runs,
    mark_docx_changes,
)


def _paragraph(doc, pieces):
    paragraph = doc.add_paragraph()
    for text, kind in pieces:
        run = paragraph.add_run(text)
        if kind == "code":
            run.font.name = "Courier New"
        elif kind == "sub":
            run.font.subscript = True
    return paragraph


def test_changed_paragraph_keeps_code_font_and_subscripts():
    doc = Document()
    paragraph = _paragraph(
        doc,
        [("Write ", None), ("H~2~O", "code"), (" and it gives H", None), ("2", "sub"), ("O in Word today", None)],
    )

    _mark_paragraph(paragraph, "Write H~2~O and it gives H2O in Word", paragraph.text)

    assert paragraph.text == "Write H~2~O and it gives H2O in Word today"
    assert [run.text for run in paragraph.runs if run.font.name == "Courier New"] == ["H~2~O"]
    assert [run.text for run in paragraph.runs if run.font.subscript] == ["2"]
    inserted = "".join(run.text for run in paragraph.runs if run.font.highlight_color == INSERTED_HIGHLIGHT)
    assert inserted.strip() == "today"


def test_swapped_words_stay_struck_through_in_place():
    doc = Document()
    paragraph = _paragraph(doc, [("cells were grown at 37 ", None), ("degrees", "code")])

    _mark_paragraph(paragraph, "cells were grown at 30 degrees", paragraph.text)

    deleted = [run for run in paragraph.runs if run.font.strike]
    assert "30" in "".join(run.text for run in deleted)
    assert deleted[0].font.color.rgb == DELETED_COLOR
    assert [run.text for run in paragraph.runs if run.font.name == "Courier New"] == ["degrees"]


def test_falls_back_to_plain_text_when_runs_miss_some_text():
    doc = Document()
    paragraph = _paragraph(doc, [("visible text", None)])

    assert _mark_runs(paragraph, "text held in an equation field", [], []) is False


def test_end_to_end_marking_preserves_formatting(tmp_path):
    old_doc = Document()
    _paragraph(old_doc, [("Concentrations use ", None), ("50 mM", "code"), (" NaCl", None)])
    old_path = tmp_path / "old.docx"
    old_doc.save(old_path)

    new_doc = Document()
    _paragraph(new_doc, [("Concentrations use ", None), ("50 mM", "code"), (" NaCl throughout", None)])
    new_path = tmp_path / "new.docx"
    new_doc.save(new_path)

    out_path, stats = mark_docx_changes(old_path, new_path, tmp_path / "marked.docx")

    assert stats["changed"] == 1
    marked = Document(out_path).paragraphs[0]
    assert [run.text for run in marked.runs if run.font.name == "Courier New"] == ["50 mM"]
    inserted = "".join(run.text for run in marked.runs if run.font.highlight_color == INSERTED_HIGHLIGHT)
    assert inserted.strip() == "throughout"
